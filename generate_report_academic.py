from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import datetime, os

BASE = r"E:\minha_java"
OUT = os.path.join(BASE, "NumberGuessingGame_Report_Academic.docx")
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2D, 0x5A, 0xA0)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def h1(t):
    h = doc.add_heading(t, level=1)
    for r in h.runs:
        r.font.color.rgb = NAVY
    return h

def h2(t):
    h = doc.add_heading(t, level=2)
    for r in h.runs:
        r.font.color.rgb = ACCENT
    return h

def h3(t):
    h = doc.add_heading(t, level=3)
    for r in h.runs:
        r.font.color.rgb = ACCENT
    return h

def para(text, bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def bullets(items):
    for it in items:
        if isinstance(it, tuple):
            p = doc.add_paragraph(style='List Bullet')
            b, rest = it
            r = p.add_run(b); r.bold = True
            p.add_run(rest)
        else:
            doc.add_paragraph(it, style='List Bullet')

def table(headers, rows, fontsize=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hh in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(hh)
        r.bold = True; r.font.size = Pt(fontsize)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(fontsize)
    doc.add_paragraph('')
    return t

def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn('w:shd'), {qn('w:fill'): 'F2F2F2', qn('w:val'): 'clear'})
    pPr.append(shd)
    return p

def placeholder_box(label, instruction):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn('w:shd'), {qn('w:fill'): 'EAF0F8', qn('w:val'): 'clear'})
    pPr.append(shd)
    r = p.add_run(f"[ {label} — PASTE APPLICATION SCREENSHOT HERE ]\n{instruction}")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ACCENT
    return p

def screenshot(path, caption):
    if not os.path.exists(path):
        placeholder_box("IMAGE MISSING", path)
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(5.4))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55,0x55,0x55)
    doc.add_paragraph('')

# ================= TITLE =================
for _ in range(2):
    doc.add_paragraph('')
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('Java Swing GUI-Based\nNumber Guessing Game'); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = NAVY
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('Academic Project Report'); r.font.size = Pt(16); r.font.color.rgb = ACCENT
i = doc.add_paragraph(); i.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = i.add_run(f"Date: {datetime.date.today().isoformat()}   •   Java Swing (Standard Core Libraries)\n"
              f"Main class: NumberGuessingGame.java   •   Range 1–100   •   10 Attempts   •   JProgressBar Tracking")
r.font.size = Pt(10)
doc.add_page_break()

# ================= 1. INDEX PAGE =================
h1('1. Index Page')
para('The following index lists all major sections and sub-sections of this report. Page numbers are provided as placeholders to be finalised after printing or PDF pagination.', italic=True, size=10)
table(['S. No.', 'Section', 'Sub-Section / Content', 'Page No.'],
[
 ['1', 'Index Page', 'Section listing with page-number placeholders', 'Pg. __ (this page)'],
 ['2', 'Project Overview', '2.1 Introduction', 'Pg. __'],
 ['', '', '2.2 Project Objectives', 'Pg. __'],
 ['', '', '2.3 Target Audience', 'Pg. __'],
 ['3', 'Technical Specifications & Libraries Used', '3.1 javax.swing.* — containers & active components', 'Pg. __'],
 ['', '', '3.2 java.awt.* & java.awt.event.* — layout, styling, events', 'Pg. __'],
 ['', '', '3.3 java.util.Random — pseudorandom engine (1–100)', 'Pg. __'],
 ['4', 'Core Application Logic & Component Flow', '4.1 Game Initialization', 'Pg. __'],
 ['', '', '4.2 Input Processing Pipeline (validation boundaries)', 'Pg. __'],
 ['', '', '4.3 Comparison Engine (Too High / Too Low / depletion)', 'Pg. __'],
 ['5', 'User Interface Interaction & Screenshot Breakdown', '5.1 Placeholder 1 — Initial State Launch Screen', 'Pg. __'],
 ['', '', '5.2 Placeholder 2 — Active Gameplay Input', 'Pg. __'],
 ['', '', '5.3 Placeholder 3 — Invalid Input / Error Boundary', 'Pg. __'],
 ['', '', '5.4 Placeholder 4 — Win State Terminal Screen', 'Pg. __'],
 ['', '', '5.5 Placeholder 5 — Loss State Terminal Screen', 'Pg. __'],
 ['6', 'Conclusion & Future Scope', '6.1 Structural stability summary', 'Pg. __'],
 ['', '', '6.2 Future evolutionary paths', 'Pg. __'],
 ['A', 'Appendix (Optional)', 'Full source listing reference', 'Pg. __'],
])
doc.add_page_break()

# ================= 2. PROJECT OVERVIEW =================
h1('2. Project Overview')
h2('2.1 Introduction')
para('The Number Guessing Game is a graphical, event-driven micro-application developed using the Java Swing graphical user interface toolkit. The application generates a concealed integer in the closed interval [1, 100]. The participant is afforded a maximum of ten (10) attempts to identify this value. After each valid submission, the system returns a directional hint — “Too High” or “Too Low” — and depletes one tracking unit from the remaining-attempts reserve.')
para('The interface is composed of a single top-level window (JFrame) partitioned into a header region, a central interaction region (input field, feedback label, textual counter, and visual progress bar), and a lower control region (Reset toggle). All interaction is mediated through a click-event ActionListener bound to the ‘Guess’ button, with supplementary lifecycle control through the ‘Reset Game’ button. Attempt consumption is reflected simultaneously in a JLabel and in a string-painted JProgressBar whose fill level and colour encode the residual game state.')
para('The pedagogical value of the project lies in its compact yet complete demonstration of state management, operational rule processing, layout configuration, and defensive input handling within a single self-contained class.')

h2('2.2 Project Objectives')
bullets([
 ('Learning event-driven UI programming: ', 'to demonstrate how a user gesture (button click) is captured by an ActionListener and routed to deterministic application logic through actionPerformed(ActionEvent).'),
 ('Practising state management: ', 'to maintain and synchronise coupled state variables (targetNumber, attemptsLeft) with multiple visual representations (JLabel, JProgressBar, button enablement) across the full game lifecycle.'),
 ('Implementing operational rule processing: ', 'to encode comparison rules (equality, less-than, greater-than), depletion rules (attemptsLeft-- per valid guess), and terminal rules (win on equality; loss on depletion to zero).'),
 ('Mastering layout configurations: ', 'to compose BorderLayout (frame), GridLayout (central stack), and FlowLayout (input and control rows) with explicit sizing, borders, fonts, and colours.'),
 ('Enforcing error boundaries for input parsing: ', 'to isolate malformed input (empty strings, non-numeric symbols, out-of-range values) via guard clauses and try-catch NumberFormatException handling without corrupting game state.'),
 ('Visualising quantitative state: ', 'to replace a purely textual counter with a JProgressBar that communicates magnitude, proportion, and urgency through fill level, embedded text, and dynamic colour.'),
])

h2('2.3 Target Audience')
para('This application is designed as an engaging educational micro-app for the following audiences:')
bullets([
 ('Introductory Java / OOP learners: ', 'a minimal, readable single-class program illustrating classes, fields, methods, constructors, and interfaces.'),
 ('HCI / GUI students: ', 'a concrete study of layout managers, component hierarchies, feedback colour semantics, and focus management.'),
 ('Instructors and evaluators: ', 'a deterministic, easily testable artefact suitable for classroom demonstration, viva examination, and laboratory assessment.'),
 ('Self-learners: ', 'a complete build-run-modify loop requiring only a standard JDK, with no external dependencies.'),
])

# ================= 3. TECHNICAL SPECIFICATIONS =================
h1('3. Technical Specifications & Libraries Used')
para('The implementation relies exclusively on standard core Java libraries shipped with the Java Development Kit (JDK). No third-party dependencies are required. The table below summarises the runtime baseline, followed by a per-library justification exactly matching the imports used in NumberGuessingGame.java.')
table(['Category', 'Specification'],
[
 ['Language', 'Java (tested on Java SE 21.0.8 LTS; compatible with Java SE 8+)'],
 ['GUI Toolkit', 'Swing (javax.swing) over AWT (java.awt)'],
 ['Top-level container', 'JFrame (400 × 300, BorderLayout, EXIT_ON_CLOSE, centred)'],
 ['Build', 'javac NumberGuessingGame.java'],
 ['Execution', 'java NumberGuessingGame (on the Event Dispatch Thread)'],
 ['External dependencies', 'None'],
])

h2('3.1 javax.swing.* — Structural Container Controls and Active Components')
para('Import: import javax.swing.*; — provides the lightweight component set from which the entire visible interface is constructed.', italic=True, size=10)
table(['Component', 'Role in This Project — Why It Was Imported'],
[
 ['JFrame', 'Top-level application window. Hosts the title ("Number Guessing Game"), size, default close operation, BorderLayout, and centre-on-screen behaviour. The game class itself extends JFrame.'],
 ['JPanel', 'Structural container for grouping. Used for centerPanel (GridLayout stack), inputPanel (FlowLayout row), and bottomPanel (Reset row); carries EmptyBorder padding.'],
 ['JLabel', 'Passive display control. Used for titleLabel (header), promptLabel ("Enter your guess:"), feedbackLabel (hint / error / terminal messages with colour semantics), and attemptsLabel (textual counter).'],
 ['JTextField (6 columns)', 'Active input control. Single-line editable text stream from which the guessed value is read via getText(). Cleared and refocused after each valid guess; setEditable(false) on termination.'],
 ['JButton ("Guess" / "Reset Game")', 'Active trigger controls. "Guess" fires the comparison pipeline; "Reset Game" invokes startNewGame(). Enablement is toggled to enforce lifecycle (Guess disabled + Reset enabled only at terminal states).'],
 ['JProgressBar (0–10)', 'Visual attempt-tracking control. Minimum 0, maximum MAX_ATTEMPTS (10), initial value full (10). setStringPainted(true) renders "N / 10 Attempts Left" inside the bar; setForeground(...) encodes urgency (green / orange / red). Smoothly depletes with attemptsLeft.'],
 ['JOptionPane (library family)', 'Member of the javax.swing.* family cited for structural completeness. In this build, terminal messaging is rendered inline via feedbackLabel (rather than a modal dialog) to keep the canvas state visible; JOptionPane remains available as a drop-in alternative for congratulatory dialog boxes without additional imports.'],
 ['BorderFactory / SwingUtilities', 'Companion Swing utilities: BorderFactory.createEmptyBorder(...) for spacing; SwingUtilities.invokeLater(...) to guarantee all construction and updates occur on the Event Dispatch Thread.'],
])

h2('3.2 java.awt.* & java.awt.event.* — Layouts, Sizing, Colours, and Event Triggers')
para('Imports: import java.awt.*; import java.awt.event.*; — provide geometry, styling, and the delegation event model underpinning interactivity.', italic=True, size=10)
table(['Element', 'Role in This Project — Why It Was Imported'],
[
 ['BorderLayout (10,10 gaps)', 'Frame-level layout: NORTH = title, CENTER = game stack, SOUTH = Reset controls.'],
 ['GridLayout (4,1,10,10)', 'Central stack layout: row 1 input, row 2 feedback, row 3 text counter, row 4 progress bar. Expanded from 3 to 4 rows to accommodate the JProgressBar.'],
 ['FlowLayout', 'Row-level layout for inputPanel (label + field + button) and bottomPanel (Reset button); natural left-to-right wrapping with gaps.'],
 ['Font (Arial BOLD/PLAIN/ITALIC)', 'Typographic hierarchy: 16pt bold title, 14pt input/feedback, 13pt counter, 12pt bold bar string.'],
 ['Color (+ BLUE/RED/ORANGE/GREEN)', 'Canvas and semantic colours: BLUE = informational, ORANGE = hint, RED = error/danger, Green (0,128,0) = success/healthy; bar foreground switches by threshold; white frame background.'],
 ['ActionListener (interface)', 'Interactive event contract. Implemented by NumberGuessingGame; actionPerformed(ActionEvent) dispatches Guess vs Reset sources.'],
 ['ActionEvent', 'Event object carrying the source control (getSource() == guessButton / restartButton) that determines which branch of game logic executes.'],
 ['Dimension-adjacent sizing', 'setSize(400,300), setLocationRelativeTo(null) for centring, JTextField column width, border insets — all AWT geometry concepts.'],
])

h2('3.3 java.util.Random — Pseudorandom Engine (Hidden Integer 1–100)')
para('Import: import java.util.Random; — supplies the pseudorandom number generator from which the concealed solution is drawn.', italic=True, size=10)
code('Random rand = new Random();\ntargetNumber = rand.nextInt(100) + 1; // uniform integer in [1, 100]')
bullets([
 ('Why Random: ', 'provides a uniform, seedable pseudorandom engine adequate for game unpredictability without external entropy sources.'),
 ('Why nextInt(100)+1: ', 'nextInt(100) yields 0–99; adding 1 shifts the bound to the specified 1–100 game range, matching the header prompt and the 1–100 range validation.'),
 ('When invoked: ', 'exactly once per lifecycle — inside startNewGame(), which is called from the constructor (initial launch) and from the Reset button (subsequent rounds) — guaranteeing a fresh concealed value each game.'),
])

# ================= 4. CORE LOGIC =================
h1('4. Core Application Logic & Component Flow')
para('This section describes the deterministic step-by-step game-loop mechanics as implemented in startNewGame(), handleGuess(), updateAttemptsBar(), endGame(), and actionPerformed().', italic=True, size=10)

h2('4.1 Game Initialization — How targetNumber and attemptsLeft Are Set on Launch')
bullets([
 ('Entry point: ', 'public static void main() delegates construction via SwingUtilities.invokeLater(() -> new NumberGuessingGame()), ensuring thread-safe creation on the Event Dispatch Thread.'),
 ('Construction: ', 'the constructor configures the JFrame (title, 400×300, EXIT_ON_CLOSE, BorderLayout), builds the header / centre / bottom panels, wires both buttons to this ActionListener, and then calls startNewGame().'),
 ('startNewGame() logic: ', 'targetNumber = new Random().nextInt(100)+1; attemptsLeft = MAX_ATTEMPTS (10); guessField cleared and setEditable(true); guessButton enabled; restartButton disabled; feedbackLabel set to "A new number has been generated. Good luck!" in BLUE; attemptsLabel set to "Attempts remaining: 10"; updateAttemptsBar() resets the bar to 10/10 full green.'),
 ('Centring: ', 'setLocationRelativeTo(null) centres the window on screen after sizing, independent of display resolution.'),
])
code('private void startNewGame() {\n'
     '    targetNumber = new Random().nextInt(100) + 1;\n'
     '    attemptsLeft = MAX_ATTEMPTS; // 10\n'
     '    guessField.setText(""); guessField.setEditable(true);\n'
     '    guessButton.setEnabled(true); restartButton.setEnabled(false);\n'
     '    feedbackLabel.setText("A new number has been generated. Good luck!");\n'
     '    attemptsLabel.setText("Attempts remaining: " + attemptsLeft);\n'
     '    updateAttemptsBar(); // 10 / 10, green\n'
     '}')

h2('4.2 Input Processing Pipeline — Text Stream, Validation Boundaries, Parsing')
para('On each ‘Guess’ click, handleGuess() executes a defensive pipeline. Invalid inputs are rejected with an explanatory red message and an early return — critically, without decrementing attemptsLeft or touching the progress bar, so the active system loop is never broken by malformed data.')
table(['Stage', 'Operation', 'Boundary / Mechanism', 'On Failure'],
[
 ['1. Read', 'String inputStr = guessField.getText().trim();', 'Trims whitespace; preserves raw stream for diagnosis', '—'],
 ['2. Empty check', 'inputStr.isEmpty()', 'Empty-field boundary', 'Red "Please enter a valid number!" + return; no cost'],
 ['3. Parse', 'Integer.parseInt(inputStr) in try block', 'Try-catch NumberFormatException boundary for non-numeric text/symbols', 'Red "Invalid characters! Enter digits only." + return; no cost'],
 ['4. Range check', 'userGuess < 1 || userGuess > 100', 'Closed-interval [1,100] boundary', 'Red "Out of range! Guess must be 1 to 100." + return; no cost'],
 ['5. Commit', 'attemptsLeft--; sync label + bar', 'Depletion step executes ONLY for validated integers', 'Proceed to Comparison Engine'],
])
code('String inputStr = guessField.getText().trim();\n'
     'if (inputStr.isEmpty()) { feedbackLabel.setText("Please enter a valid number!"); ... return; }\n'
     'try { userGuess = Integer.parseInt(inputStr); }\n'
     'catch (NumberFormatException ex) { feedbackLabel.setText("Invalid characters! Enter digits only."); ... return; }\n'
     'if (userGuess < 1 || userGuess > 100) { feedbackLabel.setText("Out of range! ..."); ... return; }')

h2('4.3 Comparison Engine — Matching, Too High / Too Low, Depletion')
para('Once input is validated, the comparison engine applies three mutually exclusive rules in order, then performs post-round housekeeping (clear field, return focus):')
table(['Rule', 'Condition', 'System Response'],
[
 ['WIN', 'userGuess == targetNumber', 'Green "Correct! You won the game!"; endGame() freezes canvas (input locked, Reset enabled). Bar retains residual value.'],
 ['LOSS', 'attemptsLeft == 0 (after decrement, no match)', 'Red "Game Over! The number was: N"; attemptsLabel "0"; bar forced to 0/10 red via updateAttemptsBar(); endGame().'],
 ['HINT (Low)', 'userGuess < targetNumber (attempts remain)', 'Orange "Too Low! Try a higher number."; label + bar already decremented and recoloured.'],
 ['HINT (High)', 'userGuess > targetNumber (attempts remain)', 'Orange "Too High! Try a lower number."; label + bar already decremented and recoloured.'],
])
code('attemptsLeft--;\n'
     'attemptsLabel.setText("Attempts remaining: " + attemptsLeft);\n'
     'updateAttemptsBar(); // value + "N / 10 Attempts Left" + green/orange/red\n'
     'if (userGuess == targetNumber) { ...win... }\n'
     'else if (attemptsLeft == 0) { ...loss, reveal target... }\n'
     'else if (userGuess < targetNumber) { ...Too Low (ORANGE)... }\n'
     'else { ...Too High (ORANGE)... }\n'
     'guessField.setText(""); guessField.requestFocus();')
para('Progress-bar colour law (updateAttemptsBar): attemptsLeft > 5 → Green (0,128,0, healthy); 3–5 → Orange (half-empty warning); ≤ 2 → Red (danger). The bar minimum is 0 and maximum is MAX_ATTEMPTS, so depletion is visually smooth and proportional.', italic=True, size=10)

# ================= 5. UI + SCREENSHOTS =================
h1('5. User Interface Interaction & Screenshot Breakdown')
para('Each sub-section below contains: (i) a structural layout table describing the UI regions, (ii) a clearly marked placeholder box indicating where the evaluator should paste the live application screenshot, (iii) the current reference screenshot embedded for convenience, and (iv) a formal explanation of button-click actions, component mutations, and operational logic.', italic=True, size=10)

def layout_table(rows):
    return table(['UI Region (Layout)', 'Component(s) & State'], rows)

# ---- 5.1 ----
h2('5.1 Screenshot Placeholder 1 — The Initial State Launch Screen')
para('Scope: startNewGame() logic, random-bound generation, label initialisation, and window centring via setLocationRelativeTo(null).', bold=True, size=10)
layout_table([
 ['NORTH (BorderLayout)', 'titleLabel: "Guess a number between 1 and 100", Arial Bold 16, centred.'],
 ['CENTER Row 1 (FlowLayout)', 'promptLabel + empty guessField + enabled Guess button.'],
 ['CENTER Row 2', 'feedbackLabel: "A new number has been generated. Good luck!" in BLUE.'],
 ['CENTER Row 3', 'attemptsLabel: "Attempts remaining: 10".'],
 ['CENTER Row 4', 'attemptsBar: 10/10 full, string-painted "10 / 10 Attempts Left", foreground GREEN.'],
 ['SOUTH (FlowLayout)', 'restartButton ("Reset Game"): DISABLED — no lifecycle override at launch.'],
 ['Frame', '400×300, EXIT_ON_CLOSE, centred via setLocationRelativeTo(null), visible on EDT.'],
])
placeholder_box('SCREENSHOT PLACEHOLDER 1', 'Paste the Initial State Launch Screen here (fresh window, full green bar, Reset disabled).')
screenshot(os.path.join(BASE,'screenshot_1_start.png'), 'Reference Fig. 5.1 — Initial state: 10 / 10 Attempts Left, green, Reset disabled.')
para('Button-click actions & operational logic:', bold=True, size=10)
bullets([
 ('No click has yet occurred. ', 'The visible state is the post-condition of startNewGame(): targetNumber = rand.nextInt(100)+1 is concealed; attemptsLeft = 10.'),
 ('Component consequences: ', 'guessField is empty and editable; Guess is armed; Reset is disarmed to prevent a redundant reset; both textual and visual counters agree at maximum.'),
 ('Why setLocationRelativeTo(null): ', 'decouples window placement from screen resolution by asking the windowing system to centre the frame, guaranteeing a consistent first impression on any display.'),
])

# ---- 5.2 ----
h2('5.2 Screenshot Placeholder 2 — Active Gameplay Input')
para('Scope: entering numbers, clicking the ‘Guess’ action button, dynamic JProgressBar drop, and conditional status-text updates.', bold=True, size=10)
layout_table([
 ['Input row', 'Numeric string typed into guessField (e.g. "25", "78"); Guess button clicked → ActionEvent → actionPerformed → handleGuess().'],
 ['Feedback row', 'ORANGE hint: "Too Low! Try a higher number." or "Too High! Try a lower number." (valid, non-terminal guess).'],
 ['Counter row', 'attemptsLabel decremented by exactly 1 per valid guess (e.g. 7, then 4).'],
 ['Progress row', 'attemptsBar value drops smoothly (7/10 green → 4/10 orange); string updates to "N / 10 Attempts Left"; colour follows threshold law.'],
 ['Control row', 'Reset remains DISABLED mid-game; focus returns to guessField via requestFocus().'],
])
placeholder_box('SCREENSHOT PLACEHOLDER 2', 'Paste Active Gameplay screenshots here (mid-game guess, partially depleted bar).')
screenshot(os.path.join(BASE,'screenshot_2_progress_green.png'), 'Reference Fig. 5.2a — Active play: guess "25", 7 / 10 green, Too Low hint.')
screenshot(os.path.join(BASE,'screenshot_3_warning_orange.png'), 'Reference Fig. 5.2b — Active play: guess "78", 4 / 10 orange, Too High hint.')
screenshot(os.path.join(BASE,'screenshot_4_danger_red.png'), 'Reference Fig. 5.2c — Active play: 2 / 10 red (danger threshold).')
para('Button-click actions & operational logic:', bold=True, size=10)
bullets([
 ('‘Guess’ click pipeline: ', 'read → validate (pass) → attemptsLeft-- → sync attemptsLabel + updateAttemptsBar() → compare → hint branch → clear field + refocus. Each click consumes exactly one tracking unit.'),
 ('Dynamic bar drop: ', 'setValue(attemptsLeft) shortens the fill proportionally (value/max); setString(...) rewrites the embedded text; setForeground(...) re-evaluates green/orange/red — all within the same EDT dispatch, so the change appears instantaneous and smooth.'),
 ('Status-text rule: ', 'directional hint colour is fixed ORANGE for non-terminal misses to distinguish guidance (orange) from errors (red) and success (green).'),
])

# ---- 5.3 ----
h2('5.3 Screenshot Placeholder 3 — Invalid Input / Error Boundary Caught')
para('Scope: visual behaviour when non-numeric text or symbols trigger the formatting exception path, highlighted in red without breaking the active loop.', bold=True, size=10)
layout_table([
 ['Input row', 'Malformed content retained in guessField (e.g. "ab@#") to aid correction; Guess clicked.'],
 ['Feedback row', 'RED error: "Invalid characters! Enter digits only." (NumberFormatException) — or "Please enter a valid number!" (empty) or "Out of range! ..." (1–100 violation).'],
 ['Counter + bar rows', 'UNTOUCHED: attemptsLabel and attemptsBar retain pre-click values (e.g. 10/10 green) — early return precedes depletion.'],
 ['Control row', 'No enablement change; game remains fully playable; focus stays available for immediate re-entry.'],
])
placeholder_box('SCREENSHOT PLACEHOLDER 3', 'Paste the Invalid Input screen here (red error message, bar unchanged).')
screenshot(os.path.join(BASE,'screenshot_7_invalid.png'), 'Reference Fig. 5.3 — Error boundary: input "ab@#", red format error, bar frozen at 10 / 10.')
para('Button-click actions & operational logic:', bold=True, size=10)
bullets([
 ('‘Guess’ with symbols: ', 'Integer.parseInt() throws NumberFormatException; the catch block sets feedbackLabel red and returns. attemptsLeft-- is never reached.'),
 ('Why red: ', 'red is reserved for non-progressing conditions (validation failures, terminal loss), contrasting with orange progress hints and green success — a consistent visual grammar.'),
 ('Loop preservation: ', 'because state mutation is sequenced strictly after validation, the system cannot be driven into a negative or desynchronised state by hostile or accidental input; the evaluator may click Guess repeatedly with invalid data and observe zero bar movement.'),
])

# ---- 5.4 ----
h2('5.4 Screenshot Placeholder 4 — The Win State Terminal Screen')
para('Scope: clicking the winning number, freezing input controls, printing congratulatory messaging, and enabling the Reset toggle.', bold=True, size=10)
layout_table([
 ['Feedback row', 'GREEN: "Correct! You won the game!" (congratulatory terminal message inline; a JOptionPane dialog may be substituted without logic change).'],
 ['Input row', 'guessField setEditable(false); Guess button setEnabled(false) — canvas frozen against further depletion.'],
 ['Counter + bar rows', 'Frozen at residual values (e.g. 6/10); no further updates possible until reset.'],
 ['Control row', 'restartButton ("Reset Game") ENABLED — sole available action; clicking it re-invokes startNewGame().'],
])
placeholder_box('SCREENSHOT PLACEHOLDER 4', 'Paste the Win State screen here (green win message, inputs disabled, Reset enabled).')
screenshot(os.path.join(BASE,'screenshot_6_win.png'), 'Reference Fig. 5.4 — Win terminal: green congratulations, locked inputs, Reset enabled.')
para('Button-click actions & operational logic:', bold=True, size=10)
bullets([
 ('Winning ‘Guess’ click: ', 'passes validation, decrements (final value retained), satisfies userGuess == targetNumber on the first branch — endGame() is invoked immediately, so the loss and hint branches are skipped.'),
 ('Freeze semantics (endGame): ', 'guessField.setEditable(false) blocks typing; guessButton.setEnabled(false) blocks events; restartButton.setEnabled(true) opens the lifecycle-override path. This tri-state toggle guarantees the terminal screen is stable and unambiguous.'),
 ('Dialog option: ', 'because javax.swing.* is already imported, JOptionPane.showMessageDialog(frame, "Correct! ...") can be added alongside the label update to produce a modal congratulatory dialog box with zero new dependencies.'),
])

# ---- 5.5 ----
h2('5.5 Screenshot Placeholder 5 — The Loss State Terminal Screen')
para('Scope: exhausting all tracking attempts, revealing the concealed solution, and opening the system lifecycle override.', bold=True, size=10)
layout_table([
 ['Feedback row', 'RED: "Game Over! The number was: N" — solution disclosure (N = targetNumber at loss time).'],
 ['Input row', 'Identical freeze to win state: non-editable field, disabled Guess.'],
 ['Counter row', 'attemptsLabel forced to "Attempts remaining: 0".'],
 ['Progress row', 'attemptsBar forced to 0/10 via updateAttemptsBar(); fill empty; foreground RED.'],
 ['Control row', 'restartButton ENABLED — lifecycle override: clicking starts a brand-new concealed number and restores 10/10 green.'],
])
placeholder_box('SCREENSHOT PLACEHOLDER 5', 'Paste the Loss State screen here (red game-over message with solution, empty red bar).')
screenshot(os.path.join(BASE,'screenshot_5_gameover.png'), 'Reference Fig. 5.5 — Loss terminal: 0 / 10 red, solution revealed, Reset enabled.')
para('Button-click actions & operational logic:', bold=True, size=10)
bullets([
 ('Final ‘Guess’ click: ', 'passes validation, decrements to exactly 0, fails the win test, satisfies attemptsLeft == 0 — the loss branch executes, revealing targetNumber that was concealed for the entire round.'),
 ('Override option: ', 'the only enabled control is Reset; its ActionEvent routes to startNewGame(), which re-randomises the bound, restores attemptsLeft = 10, resets all labels/bar/buttons — a complete lifecycle restart without JVM relaunch.'),
 ('Pedagogical note: ', 'the loss screen demonstrates graceful degradation: resource exhaustion is communicated quantitatively (0/10), textually (0 remaining), chromatically (red), and procedurally (freeze + single recovery action).'),
])

# ================= 6. CONCLUSION =================
h1('6. Conclusion & Future Scope')
h2('6.1 Summary — Structural Stability of the UI Framework')
para('The resulting application exhibits high structural stability. The single-JFrame, three-region composition (BorderLayout frame; GridLayout centre; FlowLayout rows) isolates concerns: presentation resides in labelled components, control resides in two buttons, and state resides in two integers. Because every mutation path passes through three narrow funnels — startNewGame() (birth), handleGuess() (evolution), and endGame() (termination) — the interface cannot desynchronise: the textual counter and the progress bar are updated atomically in the same method, button enablement is toggled only at lifecycle boundaries, and invalid input is quarantined before any state change.')
para('The JProgressBar augmentation materially improves usability without compromising this stability: it adds a proportional, colour-graded channel that mirrors the authoritative integer state, resets deterministically with the lifecycle, and introduces no new threads, listeners, or dependencies. The program compiles under a standard JDK, runs on any display-equipped host, and degrades gracefully under all tested input conditions.')

h2('6.2 Future Evolutionary Paths')
para('The following extensions are proposed in order of increasing complexity, each building directly on the current architecture:')
bullets([
 ('Local persistent text-file high scores: ', 'append each victory (date, attempts used, target) to a plain-text or CSV log (java.io / java.nio.file); display fewest-guesses best score in the header and add a "View Scores" control that reads the file back. Preserves the dependency-free philosophy while introducing file I/O and data persistence.'),
 ('Embedded audio feedback tracks: ', 'integrate short sampled cues via javax.sound.sampled (success chime, error buzz, click tick) triggered alongside the existing colour branches; include a mute toggle. No external libraries required; reinforces multimodal feedback design.'),
 ('Dynamic system difficulty toggles: ', 'add a difficulty selector (e.g. Easy: range 1–50 / 15 attempts; Normal: 1–100 / 10; Hard: 1–200 / 6) that re-parameterises MAX_ATTEMPTS, the Random bound, and the JProgressBar maximum at reset time; re-derive the green/orange/red thresholds proportionally. Demonstrates parameterised state and adaptive UI scaling.'),
 ('Further options: ', 'Nimbus Look-and-Feel for a modern bar skin; guess-history list (JList); range-narrowing hints ("between 20 and 40"); JUnit tests for validation and colour thresholds; JOptionPane congratulatory dialogs as a complement to inline labels.'),
])
para('Each path preserves the established event-driven loop and validation boundaries while extending the application from a classroom micro-app toward a polished, replayable desktop game.', italic=True, size=10)

doc.add_page_break()
h1('Appendix A — Build & Run Reference')
code('cd /d E:\\minha_java\njavac NumberGuessingGame.java\njava NumberGuessingGame\n\n'
     '# Regenerate screenshots (optional)\njavac GameScreenshots.java\njava GameScreenshots')
para('Main source: NumberGuessingGame.java (JFrame + ActionListener; fields targetNumber, attemptsLeft, MAX_ATTEMPTS = 10; components guessField, guessButton, restartButton, feedbackLabel, attemptsLabel, attemptsBar; methods startNewGame, updateAttemptsBar, handleGuess, endGame, actionPerformed, main). Report generated with Python python-docx; screenshots 600×420 PNG rendered from identical bar value/colour logic.', size=9, italic=True)

doc.save(OUT)
print('Saved', OUT)
