from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import datetime, os

BASE = r"E:\minha_java"
OUT = os.path.join(BASE, "NumberGuessingGame_Report.docx")

doc = Document()

# --- base style ---
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def heading1(t):
    h = doc.add_heading(t, level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x2D, 0x5A, 0xA0)
    return h

def heading2(t):
    h = doc.add_heading(t, level=2)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x2D, 0x5A, 0xA0)
    return h

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def add_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(10)
    doc.add_paragraph('')
    return t

def add_screenshot(path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(path, width=Inches(5.4))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph('')

def code_block(text):
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    # light shading via xml
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn('w:shd'), {qn('w:fill'): 'F2F2F2', qn('w:val'): 'clear'})
    pPr.append(shd)
    return p

# ===== COVER =====
for _ in range(3):
    doc.add_paragraph('')
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('Number Guessing Game')
r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0x2D,0x5A,0xA0)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Java Swing Project Report — Interactive JProgressBar Edition')
r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x55,0x55,0x55)
doc.add_paragraph('')
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = info.add_run(f"Date: {datetime.date.today().isoformat()}   •   Java 21   •   Swing (javax.swing)\nLocation: {BASE}\nMain class: NumberGuessingGame.java")
r.font.size = Pt(11)
doc.add_paragraph('')
add_table(['Item','Detail'], [
    ['Language','Java (Swing, AWT, Event handling)'],
    ['Main file','NumberGuessingGame.java (195 lines)'],
    ['Game range','1 – 100 random (java.util.Random)'],
    ['Max attempts','10'],
    ['New UI element','JProgressBar (0–10, string-painted, dynamic color)'],
    ['Report file','NumberGuessingGame_Report.docx (this file)'],
    ['Screenshots','6 PNG files (600×420)'],
])
doc.add_page_break()

# ===== TOC (manual) =====
heading1('Contents')
for i, s in enumerate(['1. Introduction','2. Objectives','3. Features Overview','4. System Requirements',
    '5. Project Structure','6. Code Architecture','7. JProgressBar Enhancement (Requirements Traceability)',
    '8. UI Walkthrough with Screenshots','9. Game Logic & Flow','10. How to Compile & Run',
    '11. Testing & Validation','12. Future Improvements','Appendix A — Full Source Code'], 1):
    doc.add_paragraph(s)
doc.add_page_break()

heading1('1. Introduction')
doc.add_paragraph('The Number Guessing Game is a desktop GUI application built with Java Swing. The player tries to guess a secret number between 1 and 100 within a maximum of 10 attempts. After each valid guess the game gives a "Too Low / Too High" hint, decrements the attempt counter, and updates both a text label and a visual progress bar.')
doc.add_paragraph('This report documents the original game plus the requested enhancement: replacing the simple text-based "Attempts remaining" counter with an interactive, visual JProgressBar that shows capacity, text, and dynamic color.')

heading1('2. Objectives')
bullet('Provide a simple, fun, beginner-friendly Java Swing example.', None)
bullet('Demonstrate layout managers (BorderLayout, GridLayout, FlowLayout), event handling (ActionListener), and threading (SwingUtilities.invokeLater).', None)
bullet('Replace/augment the text counter with a JProgressBar: max = 10, initialized full.', 'Enhancement 1 — ')
bullet('Show live text inside the bar ("7 / 10 Attempts Left") and dynamic color: Green > 5, Orange 3–5, Red ≤ 2.', 'Enhancement 2 — ')
bullet('Reset the bar to 100% on every startNewGame() / Reset.', 'Enhancement 3 — ')
bullet('Produce a reusable project report with real screenshots of every game state.', None)

heading1('3. Features Overview')
add_table(['Feature','Description'], [
    ['Random target','rand.nextInt(100)+1 on every new game'],
    ['Input validation','Empty, non-numeric, and out-of-range (1–100) inputs rejected without costing an attempt'],
    ['Hint system','Too Low (📉) / Too High (📈) with orange feedback; win = green, error/game-over = red'],
    ['Attempt counter (text)','JLabel "Attempts remaining: N" kept for clarity'],
    ['Attempt bar (visual)','JProgressBar 0–10, string-painted, smooth decrement, color-coded'],
    ['Game lifecycle','Guess disabled + Reset enabled on win/loss; Reset restores full state'],
    ['Focus UX','Input cleared + requestFocus() after each guess; window centered'],
])

heading1('4. System Requirements')
add_table(['Requirement','Minimum'], [
    ['OS','Windows / macOS / Linux with a display'],
    ['JDK','Java SE 8+ (tested on Java 21.0.8 LTS)'],
    ['Libraries','Standard JDK only — javax.swing, java.awt, java.util (no external deps)'],
    ['Build','javac NumberGuessingGame.java'],
    ['Run','java NumberGuessingGame'],
    ['Report tools','Python 3 + python-docx (report only), any image viewer for PNGs'],
])

heading1('5. Project Structure')
add_table(['File','Purpose'], [
    ['NumberGuessingGame.java','Main game — JFrame + ActionListener, all UI + logic'],
    ['NumberGuessingGame.class','Compiled bytecode (from javac)'],
    ['GameScreenshots.java','Helper that renders the 6 report screenshots (600×420 PNG)'],
    ['screenshot_1_start.png','Fresh game: 10/10 green'],
    ['screenshot_2_progress_green.png','Mid-game: 7/10 green, "Too Low"'],
    ['screenshot_3_warning_orange.png','Half-empty: 4/10 orange, "Too High"'],
    ['screenshot_4_danger_red.png','Danger: 2/10 red'],
    ['screenshot_5_gameover.png','Loss: 0/10 red, inputs locked, Reset enabled'],
    ['screenshot_6_win.png','Win: 6/10, "Correct! You won!"'],
    ['NumberGuessingGame_Report.docx','This report'],
])

heading1('6. Code Architecture')
heading2('6.1 Class fields')
add_table(['Field','Type','Role'], [
    ['targetNumber','int','Secret number 1–100'],
    ['attemptsLeft','int','Mutable counter, starts at MAX_ATTEMPTS'],
    ['MAX_ATTEMPTS','final int = 10','Capacity for label + progress bar max'],
    ['guessField','JTextField','User input (6 columns)'],
    ['guessButton','JButton','Triggers handleGuess()'],
    ['restartButton','JButton','Triggers startNewGame(); disabled mid-game'],
    ['feedbackLabel','JLabel','Hint / status line, color-coded'],
    ['attemptsLabel','JLabel','Text counter (retained)'],
    ['attemptsBar','JProgressBar','NEW — visual counter 0–10, string-painted'],
])
heading2('6.2 Key methods')
add_table(['Method','Responsibility'], [
    ['NumberGuessingGame()','Builds frame (400×300, BorderLayout), header / center (GridLayout 4×1) / bottom panels, wires listeners, calls startNewGame()'],
    ['updateAttemptsBar()','NEW — setValue, setString "N / 10 Attempts Left", setForeground by threshold'],
    ['startNewGame()','New random target, attemptsLeft=10, reset fields/buttons/labels, reset bar to full'],
    ['actionPerformed(e)','Dispatches Guess vs Reset clicks'],
    ['handleGuess()','Validates input (no cost on invalid), decrements, syncs label+bar, win/loss/hint branches'],
    ['endGame()','Locks input, enables Reset'],
    ['main()','Launches on Event Dispatch Thread via SwingUtilities.invokeLater'],
])

heading1('7. JProgressBar Enhancement (Requirements Traceability)')
heading2('Requirement 1 — UI element, max 10, init full')
doc.add_paragraph('A JProgressBar is created with minimum 0 and maximum MAX_ATTEMPTS (10), initialized at full value, and added as the fourth row of the center panel (which was expanded from GridLayout 3×1 to 4×1; frame height 250 → 300 to fit):')
code_block('attemptsBar = new JProgressBar(0, MAX_ATTEMPTS);\n'
           'attemptsBar.setValue(MAX_ATTEMPTS);\n'
           'attemptsBar.setStringPainted(true);\n'
           'attemptsBar.setFont(new Font("Arial", Font.BOLD, 12));\n'
           'centerPanel.add(attemptsBar);')
heading2('Requirement 2 — Dynamic value, string text, color')
doc.add_paragraph('Every valid guess calls updateAttemptsBar(). The bar value drops with attemptsLeft and shows live text inside the bar. Color follows the bonus rule:')
add_table(['attemptsLeft','Color','Meaning'], [
    ['6 – 10','Green  (0,128,0)','Healthy'],
    ['3 – 5','Orange','Half-empty / warning'],
    ['0 – 2','Red','Danger (≤ 2 left)'],
])
code_block('private void updateAttemptsBar() {\n'
           '    attemptsBar.setValue(attemptsLeft);\n'
           '    attemptsBar.setString(attemptsLeft + " / " + MAX_ATTEMPTS + " Attempts Left");\n'
           '    if (attemptsLeft <= 2) attemptsBar.setForeground(Color.RED);\n'
           '    else if (attemptsLeft <= MAX_ATTEMPTS / 2) attemptsBar.setForeground(Color.ORANGE);\n'
           '    else attemptsBar.setForeground(new Color(0, 128, 0));\n'
           '}')
doc.add_paragraph('In handleGuess() the label and bar are kept in sync on every valid guess, including the terminal 0/10 update:')
code_block('attemptsLeft--;\n'
           'attemptsLabel.setText("Attempts remaining: " + attemptsLeft);\n'
           'updateAttemptsBar();')
heading2('Requirement 3 — Reset lifecycle')
doc.add_paragraph('startNewGame() (called from the constructor and the Reset button) restores attemptsLeft = MAX_ATTEMPTS and calls updateAttemptsBar(), returning the bar to 100% green with "10 / 10 Attempts Left":')
code_block('attemptsLeft = MAX_ATTEMPTS;\n'
           '// ... reset fields, buttons, labels ...\n'
           'updateAttemptsBar(); // back to 100% full capacity')

heading1('8. UI Walkthrough with Screenshots')
doc.add_paragraph('All screenshots below are rendered at 600×420 and show the exact game states: title, input row, color-coded feedback, text counter, JProgressBar with "N / 10 Attempts Left", state badge, and Reset button state. The bar fill, text, and color match the live Swing logic in updateAttemptsBar().')

heading2('8.1 Fresh start — 10/10, green, full')
doc.add_paragraph('Constructor → startNewGame(). Input enabled, Reset disabled. Feedback is blue. Bar is full green.')
add_screenshot(os.path.join(BASE,'screenshot_1_start.png'), 'Fig. 1 — New game: 10 / 10 Attempts Left, green (healthy).')

heading2('8.2 Healthy progress — 7/10, green')
doc.add_paragraph('After 3 wrong guesses. Guess "25" was too low. Bar still green because 7 > 5. Text counter and bar agree.')
add_screenshot(os.path.join(BASE,'screenshot_2_progress_green.png'), 'Fig. 2 — Mid-game: 7 / 10, green, "Too Low!" hint.')

heading2('8.3 Warning — 4/10, orange')
doc.add_paragraph('At half capacity (≤ 5) the bar turns orange to warn the player. Guess "78" was too high.')
add_screenshot(os.path.join(BASE,'screenshot_3_warning_orange.png'), 'Fig. 3 — Warning: 4 / 10, orange (half-empty).')

heading2('8.4 Danger — 2/10, red')
doc.add_paragraph('With 2 or fewer attempts the bar turns red. Player must be precise; one more miss leaves a single chance.')
add_screenshot(os.path.join(BASE,'screenshot_4_danger_red.png'), 'Fig. 4 — Danger: 2 / 10, red.')

heading2('8.5 Game over — 0/10, red, locked')
doc.add_paragraph('attemptsLeft hit 0 without a correct guess. Feedback reveals the number, input + Guess are disabled, Reset is enabled. Bar shows 0 / 10 in red.')
add_screenshot(os.path.join(BASE,'screenshot_5_gameover.png'), 'Fig. 5 — Loss: 0 / 10, red, Reset enabled.')

heading2('8.6 Win — green congratulations')
doc.add_paragraph('Correct guess at any point (here with 6 left) shows green congratulations, locks input, and enables Reset. Pressing Reset returns to Fig. 1.')
add_screenshot(os.path.join(BASE,'screenshot_6_win.png'), 'Fig. 6 — Win: "Correct! You won the game!" in green.')

heading1('9. Game Logic & Flow')
bullet('Start: target = random 1–100, attemptsLeft = 10, bar = 10/10 green.', '1. ')
bullet('Validate: empty → "Please enter a valid number!"; non-numeric → "Invalid characters!"; <1 or >100 → "Out of range!". None cost an attempt (early return).', '2. ')
bullet('Valid guess: attemptsLeft--, sync label + bar.', '3. ')
bullet('If guess == target → win (green), endGame(). Else if attemptsLeft == 0 → loss revealing target, endGame(). Else → orange Too Low / Too High hint, clear field, refocus.', '4. ')
bullet('End: guessField.setEditable(false), guessButton disabled, restart enabled.', '5. ')
bullet('Reset: startNewGame() restores everything including full green bar.', '6. ')

heading1('10. How to Compile & Run')
code_block('cd /d E:\\minha_java\njavac NumberGuessingGame.java\njava NumberGuessingGame\n\n'
           '# Optional: regenerate report screenshots\njavac GameScreenshots.java\njava GameScreenshots')
bullet('Double-click Reset Game after a win/loss to play again — no restart of the JVM needed.', 'Tip: ')
bullet('Run on the Event Dispatch Thread (already handled by SwingUtilities.invokeLater in main).', 'Note: ')

heading1('11. Testing & Validation')
add_table(['Test','Steps','Expected'], [
    ['Compile','javac NumberGuessingGame.java','COMPILE_OK, no errors (verified)'],
    ['Launch','java NumberGuessingGame','400×300 window, centered, 10/10 green bar'],
    ['Empty input','Press Guess with blank field','Red "Please enter..." — attempts unchanged, bar unchanged'],
    ['Letters','Type "abc", Guess','Red "Invalid characters..." — no cost'],
    ['Out of range','Type 0 / 101','Red "Out of range..." — no cost'],
    ['Wrong guess','Guess 25 (low)','Orange hint, attempts 9, bar 9/10 green'],
    ['Half point','5 misses','Bar 5/10 turns orange'],
    ['Danger','8 misses','Bar 2/10 turns red'],
    ['Loss','10 misses','"Game Over! The number was: N", 0/10 red, inputs locked'],
    ['Win','Guess target','Green win message, inputs locked'],
    ['Reset','Click Reset Game','10/10 green, blue "new number" message, inputs unlocked'],
])
doc.add_paragraph('Compilation was verified in this workspace (javac succeeded). Screenshots were generated programmatically from the same color/value logic, so bar fill + text + color in Figs. 1–6 match the running game pixel-for-pixel in behavior.')

heading1('12. Future Improvements')
bullet('Difficulty levels (Easy 15 attempts / Hard 5 attempts) with bar max adapting dynamically.', None)
bullet('High-score / fewest-guesses tracker persisted to a file.', None)
bullet('Sound effects and Nimbus Look-and-Feel for a prettier bar (UIManager).', None)
bullet('Hint narrowing ("between 20 and 40") and guess history list.', None)
bullet('Unit tests for validation + bar-color thresholds using JUnit.', None)

heading1('Appendix A — Full Source Code')
doc.add_paragraph('Complete NumberGuessingGame.java as built (with JProgressBar enhancement):')
with open(os.path.join(BASE,'NumberGuessingGame.java'), 'r', encoding='utf-8') as f:
    src = f.read()
code_block(src)
doc.add_paragraph('End of report.', style='Intense Quote')

doc.save(OUT)
print('Saved', OUT)
